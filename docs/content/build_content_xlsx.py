import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]
TODAY = "2026-10-10"

FONT = "Arial"
F = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)
HEAD_FILL = PatternFill("solid", fgColor="1F2A44")
CUR_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CONF_FILL = PatternFill("solid", fgColor="E2EFDA")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

ACTIONS = ["Keep as is", "Replace with new content", "Remove", "Approve current text"]


def status_of(text, confirmed=False):
    if confirmed:
        return "Confirmed"
    if "[Placeholder" in str(text) or str(text).startswith("Consultant name") or text in ("Role", "Location", "Specialism"):
        return "Placeholder"
    return "Draft - needs approval"


wb = Workbook()

# ---------------- How to use ----------------
ws = wb.active
ws.title = "How to use"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 95
rows = [
    ("title", "DigitizWork website content sheet"),
    ("text", f"Current content exported from the site on {TODAY}. Fill this in and send it back; Claude updates the website from it."),
    ("blank", ""),
    ("head", "How to fill it in"),
    ("step", ("1", "Go through each tab. Column 'Current content' (grey) shows what is on the site today. Do not edit it.")),
    ("step", ("2", "To change something, type the new text in 'New content' (yellow) and set 'Action' to 'Replace with new content'.")),
    ("step", ("3", "If the current text is fine, set 'Action' to 'Approve current text'. To delete an item, choose 'Remove'.")),
    ("step", ("4", "Leave 'Action' empty (or 'Keep as is') for anything you have not decided yet.")),
    ("step", ("5", "Use 'Notes for Claude' for instructions, e.g. 'make this shorter' or 'add a 7th service like this one'.")),
    ("step", ("6", "Fill in the 'Consultants', 'Discovery questions' and 'Images' tabs where you can.")),
    ("step", ("7", "Add a line to 'Change log' each time you send the sheet back, then send the file to Claude.")),
    ("blank", ""),
    ("head", "Colour legend"),
    ("legend", ("Grey cell", "Current content on the site. Read-only reference.", CUR_FILL)),
    ("legend", ("Yellow cell", "Your input: new content, action, answers.", INPUT_FILL)),
    ("legend", ("Green cell", "Already confirmed by you. Change only if it is wrong.", CONF_FILL)),
    ("blank", ""),
    ("head", "Status column meanings"),
    ("step", ("Placeholder", "Temporary text written to show the layout. Must be replaced or removed before launch.")),
    ("step", ("Draft - needs approval", "Proposed marketing copy. Approve it, or replace it with your own.")),
    ("step", ("Confirmed", "Supplied or approved by you.")),
    ("blank", ""),
    ("head", "Example of a filled-in row"),
]
r = 1
for kind, val in rows:
    if kind == "title":
        ws.cell(r, 2, val).font = F(size=16, bold=True, color="1F2A44")
    elif kind == "text":
        ws.cell(r, 2, val).font = F(color="595959")
    elif kind == "head":
        c = ws.cell(r, 2, val)
        c.font = F(size=12, bold=True, color="1F2A44")
    elif kind == "step":
        ws.cell(r, 2, val[0]).font = F(bold=True)
        c = ws.cell(r, 3, val[1]); c.font = F(); c.alignment = WRAP
    elif kind == "legend":
        a = ws.cell(r, 2, val[0]); a.font = F(bold=True); a.fill = val[2]; a.border = BORDER
        c = ws.cell(r, 3, val[1]); c.font = F()
    r += 1

# Example row
ex_headers = ["Field ID", "Current content", "New content", "Action", "Notes for Claude"]
ex_values = ["hero.title", "Any software solution,", "AI software, built for your business.", "Replace with new content", "Keep it under 40 characters."]
for i, h in enumerate(ex_headers):
    ws.cell(r, 2, h).font = F(bold=True)
    c = ws.cell(r, 3, ex_values[i]); c.font = F(italic=True, color="595959")
    c.fill = CUR_FILL if h == "Current content" else (INPUT_FILL if h in ("New content", "Action", "Notes for Claude") else PatternFill())
    c.border = BORDER
    r += 1
r += 1
ws.cell(r, 2, "Tip").font = F(bold=True)
c = ws.cell(r, 3, "Text marked [Placeholder] is temporary. 'Field ID' tells Claude exactly where the text goes on the site, so please don't change it.")
c.font = F(); c.alignment = WRAP

# ---------------- Content sheet builder ----------------
CONTENT_HEADERS = ["Field ID", "Section", "Item", "Field", "Current content", "Status", "New content", "Action", "Characters (new)", "Guidance", "Notes for Claude"]
WIDTHS = [26, 16, 18, 18, 55, 18, 55, 22, 11, 30, 35]


def content_sheet(title, entries):
    s = wb.create_sheet(title)
    s.freeze_panes = "E2"
    for i, h in enumerate(CONTENT_HEADERS, 1):
        c = s.cell(1, i, h)
        c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
        s.column_dimensions[c.column_letter].width = WIDTHS[i - 1]
    s.row_dimensions[1].height = 30
    dv = DataValidation(type="list", formula1='"' + ",".join(ACTIONS) + '"', allow_blank=True)
    s.add_data_validation(dv)
    row = 2
    last_section = None
    for e in entries:
        fid, section, item, field, current, guidance = e[:6]
        confirmed = len(e) > 6 and e[6]
        if section != last_section:
            c = s.cell(row, 1, section.upper())
            c.font = F(bold=True, color="1F2A44")
            for col in range(1, len(CONTENT_HEADERS) + 1):
                s.cell(row, col).fill = SECTION_FILL
            row += 1
            last_section = section
        st = status_of(current, confirmed)
        vals = [fid, section, item, field, current, st, None, None, f'=IF(G{row}="","",LEN(G{row}))', guidance, None]
        for i, v in enumerate(vals, 1):
            c = s.cell(row, i, v)
            c.font = F(); c.alignment = WRAP; c.border = BORDER
        s.cell(row, 1).font = F(size=9, color="595959")
        s.cell(row, 5).fill = CONF_FILL if confirmed else CUR_FILL
        stc = s.cell(row, 6)
        stc.font = F(bold=True, color={"Placeholder": "C00000", "Confirmed": "548235"}.get(st, "BF8F00"))
        for col in (7, 8, 11):
            s.cell(row, col).fill = INPUT_FILL
        s.cell(row, 9).alignment = Alignment(horizontal="center", vertical="top")
        dv.add(f"H{row}")
        row += 1
    # spare rows for additions
    c = s.cell(row, 1, "NEW ITEMS (add rows below if you want extra content)")
    c.font = F(bold=True, color="1F2A44")
    for col in range(1, len(CONTENT_HEADERS) + 1):
        s.cell(row, col).fill = SECTION_FILL
    row += 1
    for _ in range(5):
        for col in range(1, len(CONTENT_HEADERS) + 1):
            c = s.cell(row, col); c.border = BORDER; c.font = F(); c.alignment = WRAP
        for col in (2, 3, 4, 7, 8, 10, 11):
            s.cell(row, col).fill = INPUT_FILL
        s.cell(row, 1, "new").font = F(size=9, color="595959")
        s.cell(row, 9, f'=IF(G{row}="","",LEN(G{row}))').alignment = Alignment(horizontal="center", vertical="top")
        dv.add(f"H{row}")
        row += 1
    s.auto_filter.ref = f"A1:{s.cell(1, len(CONTENT_HEADERS)).column_letter}{row - 1}"
    return s


# ---------------- Site & Contact ----------------
site_entries = [
    ("site.name", "Company", "", "Company name", "DigitizWork", "As it should appear everywhere.", True),
    ("site.description", "Company", "", "One-line description (footer)", "[Placeholder] One sentence describing what DigitizWork does and who it helps.", "Max ~140 characters."),
    ("site.contactEmail", "Contact details", "", "Public email", "consultant@digitizwork.com", "Shown in Contact and footer.", True),
    ("site.location", "Contact details", "", "Public location", "Melbourne, Australia", "City and country only.", True),
    ("site.phone", "Contact details", "", "Public phone (optional)", "", "Leave empty to show no phone."),
    ("site.address", "Contact details", "", "Street address (optional)", "", "Only if you want it public."),
    ("site.formInbox", "Contact details", "", "Inbox that receives form messages (not shown on site)", "", "Can be the same as the public email."),
    ("site.legalName", "Company", "", "Legal entity name (footer)", "", "e.g. DigitizWork Pty Ltd. Leave empty if not needed."),
    ("site.abn", "Company", "", "ABN (optional)", "", "Only if you want it in the footer."),
    ("site.linkedin", "Company", "", "Company LinkedIn URL (optional)", "", "Full URL."),
    ("nav.headerCta", "Navigation", "Header button", "Label", "Start a project", "2-3 words."),
    ("nav.services", "Navigation", "Menu item", "Label", "Services", "1-2 words."),
    ("nav.approach", "Navigation", "Menu item", "Label", "How we work", "Remove if the section is removed."),
    ("nav.whyUs", "Navigation", "Menu item", "Label", "Why us", "1-2 words."),
    ("nav.about", "Navigation", "Menu item", "Label", "About", "1-2 words."),
    ("nav.contact", "Navigation", "Menu item", "Label", "Contact", "1-2 words."),
    ("footer.legalLine", "Footer", "", "Copyright line", "DigitizWork. All rights reserved.", "Year is added automatically."),
]
content_sheet("Site & Contact", site_entries)

# ---------------- Home page ----------------
home = [
    ("hero.eyebrow", "Hero", "", "Small label above heading", "AI-powered software delivery", "2-5 words."),
    ("hero.title", "Hero", "", "Main heading (part 1)", "Any software solution,", "Heading = part 1 + part 2. Max ~40 characters each."),
    ("hero.titleHighlight", "Hero", "", "Main heading (part 2, gradient colour)", "built and delivered with AI.", "Max ~40 characters."),
    ("hero.lead", "Hero", "", "Intro paragraph", "[Placeholder] From web and mobile apps to intelligent automation and AI assistants, we design, build and run software that puts AI to work for your business, guided by experienced consultants from first idea to production.", "1-2 sentences, max ~220 characters."),
    ("hero.primaryCta", "Hero", "", "Main button", "Start a project", "2-4 words. Goes to Contact."),
    ("hero.secondaryCta", "Hero", "", "Second button", "Explore what we build", "2-4 words. Goes to Services."),
    ("hero.highlights.1", "Hero", "Highlight 1", "Short point", "[Placeholder] Idea to production", "2-4 words."),
    ("hero.highlights.2", "Hero", "Highlight 2", "Short point", "[Placeholder] Human-led, AI-accelerated", "2-4 words."),
    ("hero.highlights.3", "Hero", "Highlight 3", "Short point", "[Placeholder] Secure by design", "2-4 words."),
    ("services.eyebrow", "Services", "", "Small label", "What we build with AI", "2-5 words."),
    ("services.title", "Services", "", "Heading", "Software for every need, accelerated by AI", "Max ~60 characters."),
    ("services.intro", "Services", "", "Intro paragraph", "From a first prototype to production systems, we pair experienced engineers with modern AI tooling to deliver faster, without cutting corners.", "1-2 sentences."),
]
svc = [
    ("web-mobile", "Web and mobile apps", "Fast, accessible web platforms and mobile apps with AI features built in from day one, such as smart search, recommendations and content generation.", "Web, iOS and Android, AI features"),
    ("ai-assistants", "AI assistants and chatbots", "Assistants that answer questions, guide customers and support your team, grounded in your own documents and data.", "Customer support, Internal knowledge, LLMs"),
    ("automation", "Intelligent automation", "Remove repetitive work by connecting your tools and letting AI handle documents, triage and routine decisions, with people in the loop where it matters.", "Workflows, Document processing, Agents"),
    ("data-analytics", "Data and analytics", "Turn scattered data into dashboards, forecasts and insights your team can act on, with models that explain what they see.", "Dashboards, Forecasting, Machine learning"),
    ("cloud-integration", "Cloud and integration", "Modern cloud architecture and reliable integrations between the systems you already use, ready to scale with AI workloads.", "Cloud, APIs, Integration"),
    ("custom-software", "Custom software", "Bespoke systems for the problems off-the-shelf tools do not solve, designed around how your organisation actually works.", "Bespoke platforms, Modernisation, MVPs"),
]
for i, (sid, t, s, tags) in enumerate(svc, 1):
    item = f"Service {i}"
    home += [
        (f"services.{sid}.title", "Services", item, "Service name", t, "2-5 words. Also used in the form dropdown."),
        (f"services.{sid}.summary", "Services", item, "Description", s, "1-2 sentences, max ~180 characters."),
        (f"services.{sid}.tags", "Services", item, "Tags (comma-separated)", tags, "Up to 3 short tags."),
    ]
home += [
    ("approach.enabled", "How we work", "", "Show this section? (Yes/No)", "Yes", "Optional section."),
    ("approach.eyebrow", "How we work", "", "Small label", "How we work", "2-4 words."),
    ("approach.title", "How we work", "", "Heading", "From idea to intelligent software in four steps", "Max ~60 characters."),
    ("approach.intro", "How we work", "", "Intro paragraph", "A clear, collaborative process: you always know what is being built, why, and what comes next.", "1 sentence."),
]
steps = [
    ("Discover", "We learn your goals, users and data, and identify where AI adds real value and where it does not."),
    ("Design", "We shape the solution together: user experience, architecture and the right AI models for the job."),
    ("Build", "We deliver in short iterations with working software you can try early, tested and secured along the way."),
    ("Launch and evolve", "We release with confidence, measure what matters and keep improving the system as your needs grow."),
]
for i, (t, d) in enumerate(steps, 1):
    home += [
        (f"approach.steps.{i}.title", "How we work", f"Step {i}", "Step name", t, "1-3 words."),
        (f"approach.steps.{i}.description", "How we work", f"Step {i}", "Description", d, "1 sentence."),
    ]
home += [
    ("whyUs.eyebrow", "Why us", "", "Small label", "Why DigitizWork", "2-4 words."),
    ("whyUs.title", "Why us", "", "Heading", "Practical AI, delivered by people who care about the details", "Max ~70 characters."),
    ("whyUs.intro", "Why us", "", "Intro", "What you can expect when you work with us.", "1 sentence."),
]
why = [
    ("Consultants, not just code", "[Placeholder] You work directly with experienced consultants who understand both the business problem and the technology."),
    ("AI where it adds value", "[Placeholder] We use AI to make software smarter and delivery faster, and we are honest when a simpler solution is better."),
    ("End-to-end delivery", "[Placeholder] Strategy, design, engineering and support from one team, so nothing gets lost between handovers."),
    ("Secure and responsible", "[Placeholder] Privacy, security and responsible AI practices are built in from the start, not bolted on at the end."),
    ("Built to grow", "[Placeholder] Clean, documented architecture that is easy to extend, so your software keeps up with your organisation."),
    ("Clear communication", "[Placeholder] Regular demos and plain-language updates, so you always know where your project stands."),
]
for i, (t, d) in enumerate(why, 1):
    home += [
        (f"whyUs.items.{i}.title", "Why us", f"Point {i}", "Title", t, "2-5 words. Only claims you can stand behind."),
        (f"whyUs.items.{i}.description", "Why us", f"Point {i}", "Description", d, "1 sentence."),
    ]
home += [
    ("aboutTeaser.eyebrow", "About (home)", "", "Small label", "About us", "2-3 words."),
    ("aboutTeaser.title", "About (home)", "", "Heading", "Real consultants, amplified by AI", "Max ~50 characters."),
    ("aboutTeaser.paragraphs.1", "About (home)", "", "Paragraph 1", "[Placeholder] DigitizWork brings together consultants who build software for a living. AI makes us faster; experience makes the results dependable.", "1-2 sentences."),
    ("aboutTeaser.paragraphs.2", "About (home)", "", "Paragraph 2", "[Placeholder] One or two sentences about who DigitizWork is and why it exists, limited to facts the owner has confirmed.", "1-2 sentences. Facts only."),
    ("aboutTeaser.cta", "About (home)", "", "Button", "Meet our consultants", "2-4 words. Goes to /about."),
    ("contact.eyebrow", "Contact", "", "Small label", "Contact", "1-2 words."),
    ("contact.title", "Contact", "", "Heading", "Let's build something intelligent", "Max ~50 characters."),
    ("contact.intro", "Contact", "", "Intro paragraph", "Tell us about your idea or the problem you want to solve. A consultant will get back to you to talk through the options. [Placeholder: response time to be confirmed]", "Say how fast you reply, e.g. 'within 1 business day'."),
    ("contact.form.title", "Contact form", "", "Form heading", "Send an enquiry", "2-4 words."),
    ("contact.form.description", "Contact form", "", "Form note", "All fields are required unless marked optional.", "1 sentence."),
    ("contact.form.fields.name", "Contact form", "Field label", "Name", "Name", ""),
    ("contact.form.fields.email", "Contact form", "Field label", "Email", "Work email", ""),
    ("contact.form.fields.company", "Contact form", "Field label", "Company (optional)", "Company", ""),
    ("contact.form.fields.service", "Contact form", "Field label", "Service dropdown", "What do you need?", "Options = service names above + 'Not sure yet'."),
    ("contact.form.fields.message", "Contact form", "Field label", "Message", "Project details", ""),
    ("contact.form.fields.messageHint", "Contact form", "Field hint", "Message hint", "A few sentences about your goals, timeline and any systems involved.", "1 sentence."),
    ("contact.form.fields.consent", "Contact form", "Field label", "Consent checkbox", "[Placeholder] I agree that DigitizWork may use these details to respond to my enquiry, as described in the privacy notice.", "Should match your privacy notice."),
    ("contact.form.submitLabel", "Contact form", "", "Submit button", "Send enquiry", "2-3 words."),
]
content_sheet("Home page", home)

# ---------------- About page ----------------
about = [
    ("aboutPage.hero.eyebrow", "Hero", "", "Small label", "About us", "2-3 words."),
    ("aboutPage.hero.title", "Hero", "", "Heading (part 1)", "The people behind", "Max ~40 characters."),
    ("aboutPage.hero.titleHighlight", "Hero", "", "Heading (part 2, gradient colour)", "your AI-powered software.", "Max ~40 characters."),
    ("aboutPage.hero.lead", "Hero", "", "Intro paragraph", "[Placeholder] DigitizWork is a team of consultants who design and build software with AI. A short, factual introduction to the company goes here once confirmed by the owner.", "1-2 sentences."),
    ("aboutPage.story.eyebrow", "Our story", "", "Small label", "Our story", "2-3 words."),
    ("aboutPage.story.title", "Our story", "", "Heading", "[Placeholder] Why DigitizWork exists", "Max ~50 characters."),
    ("aboutPage.story.paragraphs.1", "Our story", "", "Paragraph 1", "[Placeholder] How and why DigitizWork started. Keep to facts the owner has confirmed: no founding year, team size or client numbers until they are provided.", "2-4 sentences."),
    ("aboutPage.story.paragraphs.2", "Our story", "", "Paragraph 2", "[Placeholder] What the company believes about building software with AI, and the kind of organisations it wants to help.", "2-4 sentences."),
    ("aboutPage.story.paragraphs.3", "Our story", "", "Paragraph 3", "[Placeholder] How the consultants work with clients day to day.", "2-4 sentences."),
    ("aboutPage.values.eyebrow", "Values", "", "Small label", "What guides us", "2-4 words."),
    ("aboutPage.values.title", "Values", "", "Heading", "How we approach every engagement", "Max ~50 characters."),
]
for i, t in enumerate(["Outcomes first", "Responsible AI", "Partnership"], 1):
    about += [
        (f"aboutPage.values.items.{i}.title", "Values", f"Value {i}", "Title", f"[Placeholder] {t}", "1-3 words."),
        (f"aboutPage.values.items.{i}.description", "Values", f"Value {i}", "Description", "[Placeholder] A short explanation of this value in practice.", "1 sentence."),
    ]
about += [
    ("aboutPage.team.eyebrow", "Consultants", "", "Small label", "Meet our consultants", "2-4 words."),
    ("aboutPage.team.title", "Consultants", "", "Heading", "The consultants you will work with", "Max ~50 characters."),
    ("aboutPage.team.intro", "Consultants", "", "Intro", "Experienced people who combine domain knowledge with hands-on AI and engineering skills.", "1 sentence. People details go in the 'Consultants' tab."),
    ("aboutPage.cta.title", "Call to action", "", "Heading", "Talk to a consultant", "2-5 words."),
    ("aboutPage.cta.text", "Call to action", "", "Text", "Tell us what you want to build. We will match you with the right people.", "1 sentence."),
    ("aboutPage.cta.button", "Call to action", "", "Button", "Start a project", "2-4 words."),
]
content_sheet("About page", about)

# ---------------- Consultants ----------------
s = wb.create_sheet("Consultants")
s.freeze_panes = "C2"
ch = ["#", "Show on site? (Yes/No)", "Full name", "Job title / role", "Short bio (2-3 sentences)", "Specialisms (comma-separated)", "Location", "LinkedIn URL", "Photo file name", "Consent to publish (Yes/No)", "Bio characters", "Notes for Claude"]
cw = [5, 12, 24, 24, 60, 32, 20, 32, 24, 14, 11, 30]
for i, h in enumerate(ch, 1):
    c = s.cell(1, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    s.column_dimensions[c.column_letter].width = cw[i - 1]
s.row_dimensions[1].height = 32
yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
s.add_data_validation(yn)
for n in range(1, 9):
    r = n + 1
    s.cell(r, 1, n)
    for col in range(1, len(ch) + 1):
        c = s.cell(r, col); c.font = F(); c.alignment = WRAP; c.border = BORDER
        if col not in (1, 11):
            c.fill = INPUT_FILL
    s.cell(r, 11, f'=IF(E{r}="","",LEN(E{r}))').alignment = Alignment(horizontal="center", vertical="top")
    yn.add(f"B{r}"); yn.add(f"J{r}")
s.cell(2, 12, "Site currently shows 4 placeholder cards. Add as many people as you need.")
r = 12
s.cell(r, 1, "Photos:").font = F(bold=True)
s.cell(r, 3, "Square, at least 800 x 800 px, JPG/PNG/WebP. Send the files with this sheet and put the file name in 'Photo file name'.").font = F()
s.cell(r + 1, 1, "Consent:").font = F(bold=True)
s.cell(r + 1, 3, "Each person must agree to their name, bio and photo being published. Cards without 'Yes' are not shown.").font = F()

# ---------------- Images ----------------
s = wb.create_sheet("Images")
s.freeze_panes = "B2"
ih = ["Image slot", "Where it appears", "Current file", "Description (alt text)", "New image file name", "Action", "Notes for Claude"]
iw = [26, 24, 34, 60, 28, 22, 34]
for i, h in enumerate(ih, 1):
    c = s.cell(1, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    s.column_dimensions[c.column_letter].width = iw[i - 1]
img_dv = DataValidation(type="list", formula1='"Keep current illustration,Replace with new image"', allow_blank=True)
s.add_data_validation(img_dv)
imgs = [
    ("heroAiCore", "Home - hero", "/images/hero/ai-core.svg", "Illustration of an AI core connected to a network of nodes, surrounded by panels showing code, a chat conversation and a data chart."),
    ("serviceWebMobile", "Home - Service 1", "/images/services/web-mobile.svg", "Illustration of a web browser window and a mobile phone showing an app interface."),
    ("serviceAiAssistants", "Home - Service 2", "/images/services/ai-assistants.svg", "Illustration of a chat conversation between a person and a glowing AI assistant."),
    ("serviceAutomation", "Home - Service 3", "/images/services/automation.svg", "Illustration of documents flowing through an AI node and coming out as completed tasks."),
    ("serviceDataAnalytics", "Home - Service 4", "/images/services/data-analytics.svg", "Illustration of a dashboard with a bar chart, a trend line and a ring chart."),
    ("serviceCloudIntegration", "Home - Service 5", "/images/services/cloud-integration.svg", "Illustration of a cloud connected to several application blocks."),
    ("serviceCustomSoftware", "Home - Service 6", "/images/services/custom-software.svg", "Illustration of layered code editor windows with code brackets."),
    ("aboutCollaboration", "Home - About section", "/images/about/collaboration.svg", "Illustration of people connected around a central AI node, representing consultants working with AI."),
    ("aboutTeamNetwork", "About page - hero", "/images/about/team-network.svg", "Illustration of a network of people linked together around a glowing core."),
    ("consultantPlaceholder", "About page - consultant without photo", "/images/team/avatar-placeholder.svg", "(decorative)"),
]
for i, row in enumerate(imgs, 2):
    for col, v in enumerate(row, 1):
        c = s.cell(i, col, v); c.font = F(); c.alignment = WRAP; c.border = BORDER
        if col in (3, 4):
            c.fill = CUR_FILL
    for col in (5, 6, 7):
        c = s.cell(i, col); c.fill = INPUT_FILL; c.border = BORDER; c.font = F(); c.alignment = WRAP
    img_dv.add(f"F{i}")
r = len(imgs) + 3
s.cell(r, 1, "Note:").font = F(bold=True)
s.cell(r, 2, "Ready-made AI image prompts and sizes for each slot are in docs/frontend/image-prompts.md. Do not use AI-generated faces for real consultants. Record the licence/terms of any image you supply.").font = F()

# ---------------- Discovery questions ----------------
s = wb.create_sheet("Discovery questions")
s.freeze_panes = "C2"
dh = ["#", "Question", "What we know so far", "Your answer", "Notes"]
dw = [7, 55, 45, 60, 30]
for i, h in enumerate(dh, 1):
    c = s.cell(1, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    s.column_dimensions[c.column_letter].width = dw[i - 1]
qs = [
    ("Q1", "What does DigitizWork do? List the services you want on the site (rough bullet points are fine).", "Message: 'DigitizWork can deliver any software solution using AI'. Service list on the site is a proposal."),
    ("Q2", "Who are your customers? Business type and size, industries, B2B or B2C.", ""),
    ("Q3", "Where do you operate and which markets do you serve? Language and spelling (Australian English?).", "Based in Melbourne, Australia (confirmed). Markets and spelling not confirmed."),
    ("Q4", "Main goal of the site? (enquiries, consultation bookings, recruiting, credibility). Which single action matters most?", "Assumed: enquiries via the contact form."),
    ("Q5", "Contact details to publish, and which inbox receives form messages?", "Email consultant@digitizwork.com and Melbourne, Australia (confirmed). Phone, address and form inbox open."),
    ("Q6", "Brand assets: logo, colours, fonts? (send files if you have them)", "None received; the current look is a proposal."),
    ("Q7", "Domain and hosting: do you own a domain? Hosting preference or budget?", ""),
    ("Q8", "Facts you can confirm: founding year, team size, founder names, legal entity name, ABN.", ""),
    ("Q9", "Competitor sites or sites you admire (2-5 links), and anything you dislike.", ""),
    ("Q10", "Privacy: confirm Australian Privacy Act applies. Cookie consent / analytics needed?", "Assumed: Australian Privacy Act; no cookies at launch."),
    ("Q11", "Analytics: OK to use privacy-friendly cookieless analytics, or none at launch?", ""),
    ("OQ-3", "Privacy notice: a small /privacy page (recommended) or a section on the page?", ""),
    ("OQ-5", "Keep the 'How we work' section? (see Home page tab, approach.enabled)", "Currently shown."),
    ("OQ-8", "Does 'launch' mean a local production build only, or a public website after testing?", ""),
]
for i, (q, text, known) in enumerate(qs, 2):
    vals = [q, text, known, None, None]
    for col, v in enumerate(vals, 1):
        c = s.cell(i, col, v); c.font = F(); c.alignment = WRAP; c.border = BORDER
    s.cell(i, 1).font = F(bold=True)
    s.cell(i, 3).fill = CUR_FILL
    s.cell(i, 4).fill = INPUT_FILL
    s.cell(i, 5).fill = INPUT_FILL

# ---------------- Change log ----------------
s = wb.create_sheet("Change log")
lh = ["Date", "Version", "Changed by", "Summary of changes", "Jira issue (filled by Claude)"]
lw = [14, 10, 20, 70, 26]
for i, h in enumerate(lh, 1):
    c = s.cell(1, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
    s.column_dimensions[c.column_letter].width = lw[i - 1]
first = [TODAY, "v1", "Claude", "Initial export of current site content (placeholder and draft copy, confirmed email and location).", ""]
for col, v in enumerate(first, 1):
    c = s.cell(2, col, v); c.font = F(); c.border = BORDER; c.alignment = WRAP
for r in range(3, 13):
    for col in range(1, 6):
        c = s.cell(r, col); c.border = BORDER; c.font = F(); c.alignment = WRAP
        if col < 5:
            c.fill = INPUT_FILL

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
